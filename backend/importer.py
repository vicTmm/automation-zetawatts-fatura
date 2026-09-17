"""Leitura do XLSX fornecido. Mantém resultados históricos sem recalculá-los."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import re
from uuid import NAMESPACE_URL, uuid5
from zipfile import BadZipFile, ZipFile

import openpyxl

from backend.calculations import INPUT_LABELS, calculate
from backend.rules import PROFILES, SHEET_RULES

INPUT_COLUMNS = {"consumption": "B", "te_supply": "C", "tusd_supply": "D", "te_injected": "E", "tusd_injected": "F", "flag": "G", "balance_credits": "K", "current_credits": "L", "enel_bill": "M", "discount_percent": "N"}
RESULT_COLUMNS = {"supply_rate": "H", "injected_rate": "I", "used_credits": "J", "total": "O", "zeta_bill": "P", "discount_total": "Q", "discount_current": "R", "discount_flag": "S", "electricity_rate": "T", "enel_bill": "M"}

LEGACY_INPUTS = {"consumption": "B", "te": "C", "tusd": "D", "flag": "E", "balance_credits": "H", "current_credits": "I", "enel_bill": "J"}
LEGACY_RESULTS = {"supply_rate": "F", "used_credits": "G", "total": "K", "zeta_bill": "L", "discount_total": "M", "discount_current": "N", "discount_flag": "O", "electricity_rate": "P", "enel_bill": "J", "consumption": "B"}
PANEL_235_INPUTS = {"enel_consumption": "K", "panel_consumption": "L", "te": "B", "tusd": "C", "flag": "D", "grid_credits": "M", "enel_bill": "N"}
PANEL_281_INPUTS = {"enel_consumption": "G", "panel_consumption": "H", "te": "C", "tusd": "D", "flag": "E", "grid_credits": "I", "enel_bill": "J"}
GD_INPUTS = {"consumption": "B", "te_supply": "C", "tusd_supply": "F", "tusd_gd1": "G", "tusd_gd2": "H", "flag": "I", "credits_gd1": "N", "credits_gd2": "O", "enel_bill": "P", "public_fees": "Q"}
GERALDO_INPUTS = {"consumption": "B", "te_supply": "C", "tusd_supply": "D", "tusd_gd1": "E", "tusd_gd2": "F", "flag": "G", "credits_gd1": "L", "credits_gd2": "M", "enel_bill": "N", "public_fees": "O"}


def read_columns(sheet, row, columns):
    return {key: number(sheet[f"{col}{row}"].value) for key, col in columns.items()}


def mapped_row(sheet, row, label, rule):
    if rule == "spec":
        inputs, results = read_columns(sheet, row, INPUT_COLUMNS), read_columns(sheet, row, RESULT_COLUMNS)
        if inputs["discount_percent"] is not None:
            inputs["discount_percent"] = format(Decimal(inputs["discount_percent"]) * 100, "f")
        results["consumption"] = inputs["consumption"]
        return inputs, results
    if rule in ("legacy_total", "essential"):
        inputs, results = read_columns(sheet, row, LEGACY_INPUTS), read_columns(sheet, row, LEGACY_RESULTS)
    elif rule in ("panel_235", "panel_281"):
        is235 = rule == "panel_235"
        inputs = read_columns(sheet, row, PANEL_235_INPUTS if is235 else PANEL_281_INPUTS)
        columns = {"consumption": "J", "supply_rate": "F", "total": "O", "zeta_bill": "P", "discount_current": "Q", "discount_flag": "R", "electricity_rate": "S", "enel_bill": "N"} if is235 else {"consumption": "B", "supply_rate": "F", "total": "K", "zeta_bill": "N", "discount_current": "L", "discount_flag": "M", "electricity_rate": "O", "enel_bill": "J"}
        results = read_columns(sheet, row, columns)
        pieces = [results.get("discount_current"), results.get("discount_flag")]
        results["discount_total"] = format(sum(Decimal(piece) for piece in pieces), "f") if all(piece is not None for piece in pieces) else None
    else:
        is_geraldo = rule == "gd_geraldo"
        inputs = read_columns(sheet, row, GERALDO_INPUTS if is_geraldo else GD_INPUTS)
        columns = {"consumption": "B", "supply_rate": "H", "used_credits": "K", "total": "P", "zeta_bill": "R", "discount_total": "V", "discount_flag": "S", "electricity_rate": "X", "enel_bill": "N", "gd1_discount": "T", "gd2_discount": "U"} if is_geraldo else {"consumption": "B", "supply_rate": "J", "used_credits": "M", "total": "R", "zeta_bill": "T", "discount_total": "X", "discount_flag": "U", "electricity_rate": "Z", "enel_bill": "P", "gd1_discount": "V", "gd2_discount": "W"}
        results = read_columns(sheet, row, columns)
        # Excel trata células vazias dos créditos como zero. Tarifa ausente só é
        # dispensável quando não existem créditos daquela modalidade.
        for category in ("gd1", "gd2"):
            if inputs["credits_" + category] is None:
                inputs["credits_" + category] = "0"
            if Decimal(inputs["credits_" + category]) == 0 and inputs["tusd_" + category] is None:
                inputs["tusd_" + category] = "0"
        pieces = [results.get("gd1_discount"), results.get("gd2_discount")]
        results["discount_current"] = format(sum(Decimal(piece or 0) for piece in pieces), "f")
    inputs["discount_percent"] = "20" if rule in ("gd", "gd_geraldo") else "10"
    return inputs, results


def number(value):
    if not isinstance(value, (int, float, Decimal)) or isinstance(value, bool):
        return None
    result = Decimal(str(value))
    return format(result, "f") if result.is_finite() else None


def import_workbook(content, store, filename="Controle de créditos ZetaGD.xlsx"):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("A planilha deve ter até 10 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 4000 or sum(x.file_size for x in archive.infolist()) > 60 * 1024 * 1024:
                raise ValueError("Planilha muito grande após descompactação.")
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=False, data_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError) as error:
        raise ValueError("Envie uma planilha .xlsx válida.") from error
    documents, report = [], []
    for sheet in workbook:
        if not sheet.title.startswith("Faturamento "):
            continue
        if sheet.max_row > 10000 or sheet.max_column > 200:
            workbook.close()
            raise ValueError("As abas de faturamento devem ter até 10 mil linhas e 200 colunas.")
        label = sheet.title.removeprefix("Faturamento ")
        cid = str(uuid5(NAMESPACE_URL, "zetawatts:" + sheet.title))
        rule = SHEET_RULES.get(label)
        if rule is None:
            report.append({"sheet": sheet.title, "rows": 0, "editable": 0, "skipped_missing": 0, "warning": "Aba nova sem mapeamento. Cadastre o cliente e selecione seu perfil na aplicação."})
            continue
        client = {"id": cid, "name": label, "legal_name": "", "document": "", "installation": "", "address": "", "discount_percent": "10", "source_sheet": sheet.title, "rule": rule}
        if label == "Icaraí":
            identifiers = str(sheet["AH4"].value or "")
            document = re.search(r"CNPJ:\s*([\d./-]+)", identifiers)
            installation = re.search(r"cliente:\s*(\d+)", identifiers, re.IGNORECASE)
            client.update(legal_name=str(sheet["AH2"].value or ""), address=str(sheet["AH3"].value or ""), document=document.group(1) if document else "", installation=installation.group(1) if installation else "", discount_percent="12.5")
        elif rule in ("gd", "gd_geraldo"):
            client["discount_percent"] = "20"
        documents.append(("client", cid, client))
        added_rows, editable_rows, missing_rows = 0, 0, 0
        for row in range(2, sheet.max_row + 1):
            reference = sheet[f"A{row}"].value
            if not isinstance(reference, (date, datetime)):
                continue
            reference = reference.strftime("%Y-%m")
            inputs, results = mapped_row(sheet, row, label, rule)
            if results.get("consumption") is None or results.get("zeta_bill") is None:
                missing_rows += 1
                continue
            complete = all(inputs.get(key) is not None for key in PROFILES[rule]["fields"])
            # Conferência não altera resultados históricos nem supõe tarifas antigas.
            editable = False
            if complete:
                try:
                    computed = calculate(inputs, rule)
                    editable = all(results.get(key) is not None and abs(Decimal(computed["results"][key]) - Decimal(results[key])) < Decimal("0.01") for key in ("zeta_bill", "discount_total", "total"))
                except ValueError:
                    editable = False
            record = {"id": cid + ":" + reference, "client_id": cid, "reference": reference, "due_date": "2026-09-20" if label == "Icaraí" and reference == "2026-08" else None, "client": client, "inputs": inputs, "results": results, "raw_discount_total": results.get("discount_total"), "line_items": computed["line_items"] if editable else [], "source": "spreadsheet", "source_location": f"{filename} / {sheet.title} / linha {row}", "rule_version": computed["rule_version"] if editable else "historical", "editable": editable, "flag_label": "Amarela" if label == "Icaraí" and reference == "2026-08" else "", "notes": "", "payment_image": None, "warnings": []}
            documents.append(("invoice", record["id"], record))
            added_rows += 1
            editable_rows += int(editable)
        report.append({"sheet": sheet.title, "rows": added_rows, "editable": editable_rows, "skipped_missing": missing_rows})
    workbook.close()
    if not report:
        raise ValueError("Nenhuma aba 'Faturamento ...' foi encontrada nesta planilha.")
    return {**store.import_documents(documents), "sheets": report}
